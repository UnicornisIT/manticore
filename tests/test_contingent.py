import io
import json
import re
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from docx.oxml.ns import qn

from contingent import group_kind, map_groups, normalize_group, parse_docx, reconcile, student_row


FIXTURES = Path(__file__).parent / 'fixtures' / 'contingent'


def document(groups):
    return {'groups': [dict(id=str(i), source_name=name, normalized_name=normalize_group(name),
                           specialty='', is_legacy=group_kind(name) == 'legacy',
                           students=[student_row(fio, str(j)) for j, fio in enumerate(names)])
                       for i, (name, names) in enumerate(groups)]}


class ContingentEngineTests(unittest.TestCase):
    def test_real_structure_fixtures(self):
        summary = parse_docx((FIXTURES / 'department-2026-2027.docx').read_bytes(), 'department-2026-2027.docx')
        self.assertEqual(summary['document_type'], 'DEPARTMENT_CONTINGENT_TABLES')
        self.assertEqual(len(summary['groups']), 32)
        self.assertEqual(len(summary['specialties']), 7)
        active = [g for g in summary['groups'] if not g['is_legacy']]
        self.assertEqual(len(active), 8)
        self.assertEqual(sum(len(g['students']) for g in active), 91)
        self.assertTrue(summary['metadata_conflict'])
        self.assertTrue(all(s['valid'] for g in summary['groups'] for s in g['students']))
        single = parse_docx((FIXTURES / 'single.docx').read_bytes(), 'single.docx')
        self.assertEqual(single['document_type'], 'SINGLE_GROUP_CONTINGENT')
        self.assertEqual(single['groups'][0]['normalized_name'], '26ЛД-9И')
        self.assertEqual(len(single['groups'][0]['students']), 2)
        self.assertEqual(single['education_period'], '2026–2030')
        fixture_doc = Document(FIXTURES / 'single.docx')
        self.assertEqual(sum(bool(p._p.xpath('./w:pPr/w:numPr')) for p in fixture_doc.paragraphs), 2)
        definitions = fixture_doc.part.numbering_part.element
        numbered_definition = next(a for a in definitions.findall(qn('w:abstractNum')) if a.get(qn('w:abstractNumId')) == '8')
        self.assertEqual(numbered_definition.find(qn('w:lvl')).find(qn('w:numFmt')).get(qn('w:val')), 'decimal')

    def test_note_normalization(self):
        row = student_row('1. Тестов\u00a0 Иван   Иванович пр.от 22.04.2026 №2204-У', 'p1')
        self.assertEqual(row['normalized_full_name'], 'тестов иван иванович')
        self.assertEqual(row['source_note'], 'пр.от 22.04.2026 №2204-У')
        self.assertEqual(row['order_date'], '2026-04-22')
        self.assertEqual(row['order_number'], '2204-У')
        self.assertIn('\u00a0', row['raw_text'])

    def test_arbitrary_notes_after_fio(self):
        for note in ['(зач.приказом № 1808-У от 18.08.2026)', 'документы предоставлены',
                     'ЛЮБОЙ ДОПОЛНИТЕЛЬНЫЙ ТЕКСТ', '— примечание', ', уточнить сведения']:
            with self.subTest(note=note):
                row = student_row('1. Тестов Иван Иванович ' + note, 'p1')
                self.assertTrue(row['valid'])
                self.assertEqual(row['normalized_full_name'], 'тестов иван иванович')
                self.assertEqual(row['source_note'], note)
        row = student_row('Тестов Иван Иванович(зач.приказом № 1808-У от 18.08.2026)', 'p1')
        self.assertEqual(row['order_number'], '1808-У')
        self.assertEqual(row['order_date'], '2026-08-18')
        self.assertEqual(row['normalized_full_name'], 'тестов иван иванович')

    def test_compound_names_and_patronymic_suffix(self):
        for name in ['Тестова-Примерова Анна-Мария Ивановна', 'Тестов Иван', 'Тестов Али Мамед оглы']:
            row = student_row(name, 'p1')
            self.assertTrue(row['valid'])
            self.assertEqual(row['normalized_full_name'], name.casefold())
        self.assertFalse(student_row('Неразборчиво ???', 'p1')['valid'])

    def test_group_prefix_and_notes_in_single_group_document(self):
        for title in ['ГРУППА 26ФМ -  11 И', 'Группа № 26ФМ-11И', 'Группа: 26ФМ-11И']:
            doc = Document()
            for text in [title, '2026-2027 учебный год', 'Начало обучения: осень 2026 года',
                         'Срок обучения – 1 год 10 месяцев', 'ВЫПУСК - июнь 2028 года',
                         'Тестов Иван Иванович (зач.приказом № 1808-У от 18.08.2026)']:
                doc.add_paragraph(text)
            stream = io.BytesIO(); doc.save(stream)
            parsed = parse_docx(stream.getvalue(), 'single-group.docx')
            group = parsed['groups'][0]
            self.assertEqual(group['normalized_name'], '26ФМ-11И')
            self.assertEqual(len(group['students']), 1)
            mapped = map_groups(parsed, [{'name': '26ФМ-11И-1'}])
            rows = reconcile(mapped, [dict(username='test', fio='Тестов Иван Иванович', group='26ФМ-11И-1')])
            self.assertEqual(rows[0]['status'], 'MATCHED')

    def test_repeated_heading_does_not_create_empty_group(self):
        doc = Document()
        for text in ['26СтО -11', 'ГРУППА 26СтО - 11', '2026-2027 учебный год',
                     'Тестов Иван Иванович', '26СтО -11', 'Примеров Петр Петрович']:
            doc.add_paragraph(text)
        stream = io.BytesIO(); doc.save(stream)
        groups = parse_docx(stream.getvalue(), 'repeated-heading.docx')['groups']
        self.assertEqual(len(groups), 2)
        self.assertEqual([len(group['students']) for group in groups], [1, 1])
        self.assertEqual([group['id'] for group in groups], ['0', '1'])

    def test_short_group_and_legacy_policy(self):
        source = document([('26ЛД - 9 И', ['Тестов Иван'])])
        groups = [{'name': '26ЛД-9И-1'}]
        self.assertEqual(map_groups(source, groups)[0]['mapping_status'], 'AUTO_NORMALIZED_GROUP')
        groups.append({'name': '26ЛД-9И-2'})
        self.assertEqual(map_groups(source, groups)[0]['target'], '')
        self.assertEqual(map_groups(source, groups, {'0': '26ЛД-9И-2'})[0]['target'], '26ЛД-9И-2')
        for name in ['211', '212', '311', '2ЛД', '3 ЛД', '2ФМ']:
            self.assertEqual(group_kind(name), 'legacy')
        self.assertEqual(group_kind('27СД-11-1'), 'modern')
        self.assertIsNone(group_kind('неизвестная группа'))

    def test_two_way_matching_duplicates_and_manual_identity(self):
        source = document([('26ЛД-9-1', ['Тестов Альфа', 'Тестов Бета', 'Тестов Гамма'])])
        groups = [{'name': '26ЛД-9-1'}]
        students = [dict(username=name, fio='Тестов ' + name, group='26ЛД-9-1') for name in ['Альфа', 'Бета', 'Дельта']]
        rows = reconcile(map_groups(source, groups), students)
        self.assertEqual([r['status'] for r in rows], ['MATCHED', 'MATCHED', 'NOT_FOUND', 'EXTRA_IN_MANTICORE'])
        students[0]['group'] = '26ЛД-9-2'
        self.assertEqual(reconcile(map_groups(source, groups), students)[0]['status'], 'WRONG_GROUP')
        students.append(dict(students[0], username='second'))
        self.assertEqual(reconcile(map_groups(source, groups), students)[0]['status'], 'AMBIGUOUS_MATCH')
        self.assertEqual(reconcile(map_groups(source, groups), students, {'0:0': 'second'})[0]['username'], 'second')
        source['groups'][0]['students'].append(source['groups'][0]['students'][0])
        self.assertEqual(reconcile(map_groups(source, groups), students)[0]['status'], 'DUPLICATE_IN_SOURCE')
        source['groups'].append(dict(source['groups'][0], id='1', normalized_name='26ЛД-9-2'))
        self.assertEqual(reconcile(map_groups(source, groups), students)[0]['status'], 'MULTIPLE_SOURCE_GROUPS')

    def test_legacy_and_specialty_are_safe(self):
        source = document([('211', ['Тестов Иван']), ('26ЛД-9-1', ['Примеров Петр'])])
        source['groups'][1]['specialty'] = 'СЕСТРИНСКОЕ ДЕЛО'
        mapped = map_groups(source, [{'name': '26ЛД-9-1'}])
        self.assertEqual(mapped[0]['mapping_status'], 'LEGACY_IGNORED')
        self.assertEqual(mapped[1]['target'], '')
        self.assertEqual(len(reconcile(mapped, [])), 1)

    def test_invalid_docx_is_friendly(self):
        for content in [b'invalid', b'']:
            with self.assertRaises(ValueError):
                parse_docx(content, 'wrong.docx')
        output = io.BytesIO(); Document().save(output)
        with self.assertRaisesRegex(ValueError, 'структуру'):
            parse_docx(output.getvalue(), 'empty.docx')

    def test_merged_cells_and_repeated_group_sections(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = '26ЛД-9-1'
        table.cell(1, 0).text = '1.'; table.cell(1, 1).text = 'Тестов Иван'
        table.cell(2, 0).text = '2.'; table.cell(2, 1).text = 'Неразборчиво ???'
        stream = io.BytesIO(); doc.save(stream)
        parsed = parse_docx(stream.getvalue(), 'merged.docx')
        self.assertEqual(len(parsed['groups']), 1)
        rows = reconcile(map_groups(parsed, [{'name': '26ЛД-9-1'}]), [])
        self.assertEqual([r['status'] for r in rows], ['NOT_FOUND', 'INVALID_ROW'])


class ContingentWorkflowTests(unittest.TestCase):
    def setUp(self):
        # Reuse existing initialized test DB, but never the production database.
        import test_app
        self.m = test_app.manticore
        self.tmp = tempfile.TemporaryDirectory()
        self.patch = mock.patch.dict(self.m.app.config, UPLOAD_FOLDER=self.tmp.name, TESTING=True)
        self.patch.start()
        self.db = Path(self.tmp.name) / 'test.db'
        with sqlite3.connect(self.m.DB_PATH) as source, sqlite3.connect(self.db) as target:
            source.backup(target)
        self.dbpatch = mock.patch.object(self.m, 'DB_PATH', str(self.db)); self.dbpatch.start()
        with sqlite3.connect(self.db) as conn:
            conn.execute('DELETE FROM students'); conn.execute('DELETE FROM groups')
            conn.execute('DELETE FROM audit_logs'); conn.execute('DELETE FROM student_group_transfers')
            conn.executemany('INSERT INTO groups(name, group_year) VALUES (?, ?)', [('26ЛД-9И-1', '2026'), ('26ЛД-9И-2', '2026')])
            conn.execute("INSERT INTO students(username, firstname, lastname, cohort1) VALUES ('moving', 'Иван Иванович', 'Тестов', '26ЛД-9И-2')")
        self.client = self.m.app.test_client()
        with self.client.session_transaction() as session:
            session['user'] = 'admin'; session['role'] = 'admin'; session['csrf_token'] = 'test-token'
        self.m.save_login_generation_settings(self.m.get_default_login_generation_rules(), setup_completed=True, updated_by='test')

    def tearDown(self):
        self.dbpatch.stop(); self.patch.stop()
        import gc
        gc.collect()
        self.tmp.cleanup()

    def post(self, path, data=None):
        return self.client.post(path, data=dict(csrf_token='test-token', **(data or {})))

    def upload(self):
        doc = Document(); doc.add_paragraph('26ЛД-9И-1'); doc.add_paragraph('Тестов Иван Иванович', 'List Number')
        content = io.BytesIO(); doc.save(content); content.seek(0)
        response = self.post('/file_work/contingent/upload', {'file': (content, 'test.docx')})
        self.assertEqual(response.status_code, 303, response.get_data(as_text=True))
        self.path = response.location
        self.assertEqual(self.client.get(self.path).status_code, 200)
        return self.path

    def state(self):
        return json.loads(next((Path(self.tmp.name) / 'contingent_previews').glob('*.json')).read_text(encoding='utf-8'))

    def test_upload_reconcile_prepare_apply_recheck_export(self):
        with sqlite3.connect(self.db) as conn: before = list(conn.iterdump())
        path = self.upload()
        self.assertEqual(self.post(path, {'year': '2026'}).status_code, 303)
        self.assertEqual(self.state()['rows'][0]['status'], 'WRONG_GROUP')
        self.assertEqual(self.post(path + '/prepare', {'selected': '0:0'}).status_code, 200)
        with sqlite3.connect(self.db) as conn: self.assertEqual(before, list(conn.iterdump()))
        token = self.state()['confirmation']
        self.assertEqual(self.post(path + '/apply', {'confirmation': token}).status_code, 303)
        self.assertEqual(self.state()['rows'][0]['status'], 'MATCHED')
        with sqlite3.connect(self.db) as conn:
            self.assertEqual(conn.execute('SELECT count(*) FROM student_group_transfers').fetchone()[0], 1)
            audit = str(conn.execute('SELECT * FROM audit_logs').fetchall())
            self.assertIn('test.docx', audit)
        self.assertEqual(self.post(path + '/apply', {'confirmation': token}).status_code, 409)
        exported = self.client.get(path + '/export')
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(exported.data))
        self.assertEqual(workbook.active.cell(2, 6).value, 'Совпадает')
        self.assertEqual(self.client.get(path).status_code, 200)

    def test_permissions_csrf_and_stale_source(self):
        path = self.upload(); self.post(path, {'year': '2026'})
        self.post(path + '/prepare', {'selected': '0:0'})
        token = self.state()['confirmation']
        with sqlite3.connect(self.db) as conn: conn.execute("UPDATE students SET cohort1='26ЛД-9И-1'")
        self.post(path + '/apply', {'confirmation': token})
        self.assertFalse(self.state()['outcomes'][0]['ok'])
        self.assertIn('изменились', self.state()['outcomes'][0]['message'])
        with sqlite3.connect(self.db) as conn:
            conn.execute("UPDATE users SET role='assistant' WHERE username='admin'")
            conn.execute("INSERT OR REPLACE INTO users(username, role, approved) VALUES ('other', 'assistant', 1)")
        with self.client.session_transaction() as session: session['role'] = 'assistant'
        self.assertEqual(self.post(path + '/prepare', {'selected': '0:0'}).status_code, 302)
        self.assertEqual(self.post(path + '/apply', {'confirmation': token}).status_code, 302)
        with self.client.session_transaction() as session: session['user'] = 'other'
        self.assertEqual(self.client.get(path).status_code, 403)
        self.assertEqual(self.client.post(path).status_code, 303)

    def test_capacity_override_is_explicit(self):
        with sqlite3.connect(self.db) as conn:
            conn.executemany('INSERT INTO students(username, cohort1) VALUES (?, ?)', [(f'full{i}', '26ЛД-9И-1') for i in range(25)])
        path = self.upload(); self.post(path, {'year': '2026'})
        self.post(path + '/prepare', {'selected': '0:0'})
        self.assertEqual(self.state()['prepared'][0]['capacity_state'], 'full')
        self.post(path + '/apply', {'confirmation': self.state()['confirmation']})
        self.assertFalse(self.state()['outcomes'][0]['ok'])
        self.post(path + '/prepare', {'selected': '0:0'})
        self.post(path + '/apply', {'confirmation': self.state()['confirmation'], 'capacity_override': 'yes'})
        self.assertTrue(self.state()['outcomes'][0]['ok'])

    def test_metadata_requires_explicit_context_and_read_only_owner(self):
        data = {'file': (io.BytesIO((FIXTURES / 'department-2026-2027.docx').read_bytes()), 'department-2026-2027.docx')}
        response = self.post('/file_work/contingent/upload', data)
        path = response.location
        self.assertTrue(self.state()['document']['metadata_conflict'])
        self.assertEqual(self.post(path, {'year': '2026'}).status_code, 200)
        self.assertNotIn('rows', self.state())
        self.assertEqual(self.post(path, {'year': '2026', 'metadata_confirmed': 'yes'}).status_code, 303)
        self.assertTrue(self.state()['rows'])
        self.assertEqual(self.client.get(path + '/export?discrepancies=1').status_code, 200)

    def test_group_status_layout_all_capacity_variants(self):
        with sqlite3.connect(self.db) as conn:
            conn.executemany('INSERT INTO students(username, cohort1) VALUES (?, ?)',
                [(f'full{i}', '26ЛД-9И-1') for i in range(25)] + [(f'over{i}', '26ЛД-9И-2') for i in range(25)])
        body = self.client.get('/add_group').get_data(as_text=True)
        self.assertIn('>Заполнена</span>', body)
        self.assertIn('>Переполнена</span>', body)
        self.assertLess(body.index('>Удалить</button>'), body.index('>Заполнена</span>'))
